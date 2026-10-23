"""Office files, spreadsheets, slides, e-books, Markdown / HTML, data, SVG, archives, subtitles."""
import json
import tarfile
import zipfile
from pathlib import Path

import pytest

from dragshift import convert, documents, media
from dragshift.util import UserError, category

needs_lo = pytest.mark.skipif(not documents.lo_available(), reason="LibreOffice not installed")
needs_ffmpeg = pytest.mark.skipif(not media.available(), reason="FFmpeg not available")


@pytest.fixture
def no_lo(monkeypatch):
    """Pretend LibreOffice isn't installed: the pure-Python paths must still work."""
    monkeypatch.setattr(documents, "find_soffice", lambda: None)


def run(src, target):
    return convert.convert(src, category(src), target)


@pytest.fixture
def csv_file(tmp_path):
    p = tmp_path / "People.csv"
    p.write_text('name,age,city\nAna,31,Lisboa\nBruno,0042,Porto\n"Silva, J",28.5,Faro\n', encoding="utf-8")
    return p


@pytest.fixture
def xlsx(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Item", "Qty", "Price"])
    ws.append(["Pen", 3, 1.5])
    ws.append(["Book", 1, 12.0])
    ws2 = wb.create_sheet("Other")
    ws2.append(["k", "v"])
    ws2.append(["a", 1])
    p = tmp_path / "Book.xlsx"
    wb.save(p)
    return p


@pytest.fixture
def pptx(tmp_path):
    from pptx import Presentation
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Hello deck"
    s.placeholders[1].text = "First point\nSecond point"
    prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = "Slide two"
    p = tmp_path / "Deck.pptx"
    prs.save(p)
    return p


@pytest.fixture
def md(tmp_path):
    p = tmp_path / "Notes.md"
    p.write_text("# Title\n\nSome **bold** text and a [link](http://x.y).\n\n- one\n- two\n\n"
                 "| a | b |\n|---|---|\n| 1 | 2 |\n\n```\nprint('hi')\n```\n", encoding="utf-8")
    return p


@pytest.fixture
def zipf(tmp_path):
    p = tmp_path / "Pack.zip"
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("a.txt", "hello")
        z.writestr("dir/b.txt", "world" * 1000)
    return p


def test_categories(tmp_path):
    for name, cat in {"a.csv": "sheet", "a.json": "sheet", "a.xlsx": "sheet", "a.md": "markdown",
                      "a.html": "html", "a.pptx": "slides", "a.odt": "word", "a.rtf": "word",
                      "a.epub": "ebook", "a.svg": "svg", "a.zip": "archive", "a.tar.gz": "archive",
                      "a.srt": "subtitle", "a.txt": "text", "a.log": "text", "a.docx": "docx"}.items():
        assert category(tmp_path / name) == cat, name


# ---------------------------------------------------------------- spreadsheets and data
def test_csv_conversions(csv_file):
    j = json.loads(run(csv_file, "JSON")[0].read_text(encoding="utf-8"))
    assert j[0] == {"name": "Ana", "age": 31, "city": "Lisboa"}
    assert j[1]["age"] == "0042"          # leading zeros are not numbers
    assert j[2] == {"name": "Silva, J", "age": 28.5, "city": "Faro"}
    import openpyxl
    ws = openpyxl.load_workbook(run(csv_file, "XLSX")[0]).active
    assert ws["B2"].value == 31 and ws["A4"].value == "Silva, J"
    assert "| name | age | city |" in run(csv_file, "MD")[0].read_text(encoding="utf-8")
    assert "<td>Lisboa</td>" in run(csv_file, "HTML")[0].read_text(encoding="utf-8")
    assert run(csv_file, "TSV")[0].read_text("utf-8-sig").splitlines()[1] == "Ana\t31\tLisboa"


def test_xlsx_sheets_become_separate_files(xlsx):
    outs = run(xlsx, "CSV")
    assert [o.name for o in outs] == ["Book Sales.csv", "Book Other.csv"]
    assert outs[0].read_text("utf-8-sig").splitlines() == ["Item,Qty,Price", "Pen,3,1.5", "Book,1,12"]
    assert len(run(xlsx, "JSON")) == 2
    assert len(run(xlsx, "HTML")) == 1


def test_sheet_to_pdf_and_docx(xlsx):
    import pymupdf
    pdf = run(xlsx, "PDF")[0]
    with pymupdf.open(pdf) as d:
        assert d.page_count == 2 and "Pen" in d[0].get_text()
    from docx import Document
    d = Document(run(xlsx, "DOCX")[0])
    assert len(d.tables) == 2 and d.tables[0].cell(1, 0).text == "Pen"


def test_json_shapes(tmp_path):
    p = tmp_path / "x.json"
    p.write_text(json.dumps([{"id": 1, "tags": ["a"]}, {"id": 2, "extra": "é"}]))
    rows = documents.load_sheets(p)[0][1]
    assert rows[0] == ["id", "tags", "extra"] and rows[1][1] == '["a"]' and rows[2][2] == "é"
    p.write_text('"just a string"')
    with pytest.raises(UserError, match="table"):
        run(p, "CSV")
    p.write_text("{nope")
    with pytest.raises(UserError, match="valid JSON"):
        run(p, "CSV")


def test_semicolon_csv_is_sniffed(tmp_path):
    p = tmp_path / "eu.csv"
    p.write_text("a;b\n1;2\n", encoding="utf-8")
    assert documents.load_sheets(p)[0][1] == [["a", "b"], ["1", "2"]]


def test_never_overwrites(csv_file):
    a, b = run(csv_file, "JSON")[0], run(csv_file, "JSON")[0]
    assert a != b and a.exists() and b.exists()


# ---------------------------------------------------------------- Markdown and HTML
def test_markdown_conversions(md):
    html = run(md, "HTML")[0].read_text(encoding="utf-8")
    assert "<h1>Title</h1>" in html and "<strong>bold</strong>" in html and "<table>" in html
    txt = run(md, "TXT")[0].read_text(encoding="utf-8")
    assert "Title" in txt and "• one" in txt and "print('hi')" in txt and "**" not in txt
    from docx import Document
    d = Document(run(md, "DOCX")[0])
    assert d.paragraphs[0].style.name == "Heading 1" and len(d.tables) == 1
    assert any(p.style.name == "List Bullet" for p in d.paragraphs)
    import pymupdf
    with pymupdf.open(run(md, "PDF")[0]) as pdf:
        assert "Title" in pdf[0].get_text()


def test_html_to_markdown_and_text(tmp_path):
    p = tmp_path / "Page.html"
    p.write_text("<html><head><title>T</title><style>p{}</style><script>var x=1</script></head><body>"
                 "<h1>Hello</h1><p>Para <b>bold</b> <a href='http://z'>z</a></p>"
                 "<ul><li>x</li><li>y</li></ul>"
                 "<table><tr><th>h1</th><th>h2</th></tr><tr><td>1</td><td>2</td></tr></table></body></html>")
    out = run(p, "MD")[0].read_text(encoding="utf-8")
    assert "# Hello" in out and "**bold**" in out and "[z](http://z)" in out and "- x" in out
    assert "| h1 | h2 |" in out and "var x" not in out and "p{}" not in out
    assert "Hello" in run(p, "TXT")[0].read_text(encoding="utf-8")
    assert run(p, "PDF")[0].exists() and run(p, "DOCX")[0].exists()


def test_plain_text_to_html(tmp_path):
    p = tmp_path / "n.txt"
    p.write_text("one <two>\n\nthree")
    assert "one &lt;two&gt;" in convert.convert(p, "text", "HTML")[0].read_text(encoding="utf-8")


# ---------------------------------------------------------------- slides
def test_pptx_text(pptx):
    txt = run(pptx, "TXT")[0].read_text(encoding="utf-8")
    assert "Slide 1: Hello deck" in txt and "• First point" in txt and "Slide 2: Slide two" in txt


def test_pptx_pdf_and_images_without_libreoffice(pptx, no_lo):
    import pymupdf
    assert "ODP" not in convert.targets_for("slides", ".pptx")
    with pymupdf.open(run(pptx, "PDF")[0]) as d:
        assert d.page_count == 2 and "Hello deck" in d[0].get_text()
    pngs = run(pptx, "PNG")
    assert [p.name for p in pngs] == ["Deck slide 1.png", "Deck slide 2.png"]


@needs_lo
def test_pptx_with_libreoffice(pptx):
    import pymupdf
    with pymupdf.open(run(pptx, "PDF")[0]) as d:
        assert d.page_count == 2
    odp = run(pptx, "ODP")[0]
    assert category(odp) == "slides" and "Hello deck" in run(odp, "TXT")[0].read_text(encoding="utf-8")


# ---------------------------------------------------------------- word-like files
def test_odt_and_rtf_without_libreoffice(tmp_path, no_lo):
    odt = tmp_path / "a.odt"
    with zipfile.ZipFile(odt, "w") as z:
        z.writestr("content.xml", '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:'
                   'xmlns:office:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">'
                   '<office:body><office:text><text:h>Heading</text:h><text:p>Body <text:span>text</text:span>'
                   '</text:p></office:text></office:body></office:document-content>')
    assert run(odt, "TXT")[0].read_text(encoding="utf-8") == "Heading\nBody text"
    rtf = tmp_path / "b.rtf"
    rtf.write_text(r"{\rtf1\ansi Hello {\b bold} world\par Second line}")
    assert "Hello bold world" in run(rtf, "TXT")[0].read_text(encoding="utf-8")
    from docx import Document
    assert any("Second line" in p.text for p in Document(run(rtf, "DOCX")[0]).paragraphs)
    assert run(rtf, "PDF")[0].exists()
    assert convert.targets_for("word", ".doc") == []     # .doc needs LibreOffice
    assert documents.missing_tool(".doc") == "LibreOffice"


@needs_lo
def test_word_formats_with_libreoffice(md):
    docx = run(md, "DOCX")[0]
    for target in ("ODT", "RTF", "DOC"):
        out = run(docx, target)[0]
        assert category(out) == "word"
        assert "Title" in run(out, "TXT")[0].read_text(encoding="utf-8")
    doc = run(docx, "DOC")[0]
    assert "DOCX" in convert.targets_for("word", ".doc") and run(doc, "PDF")[0].exists()
    assert "Title" in run(docx, "HTML")[0].read_text(encoding="utf-8")
    assert "Title" in run(docx, "MD")[0].read_text(encoding="utf-8")


@needs_lo
def test_xls_and_ods(csv_file):
    ods, xls = run(csv_file, "ODS")[0], run(csv_file, "XLS")[0]
    for f in (ods, xls):
        assert category(f) == "sheet"
        assert "Lisboa" in run(f, "CSV")[0].read_text("utf-8-sig")


# ---------------------------------------------------------------- e-books and SVG
def test_ebook_and_svg(tmp_path):
    import pymupdf
    book = pymupdf.open()
    page = book.new_page()
    page.insert_text((72, 72), "Chapter one", fontsize=20)
    pdf = tmp_path / "Book.pdf"
    book.save(pdf)
    # PyMuPDF can write XPS-free page documents; a CBZ is just images in a zip
    from PIL import Image
    cbz = tmp_path / "Comic.cbz"
    with zipfile.ZipFile(cbz, "w") as z:
        for i in range(2):
            Image.new("RGB", (120, 160), (200, 50 * i, 80)).save(tmp_path / f"{i}.png")
            z.write(tmp_path / f"{i}.png", f"{i}.png")
    assert convert.targets_for("ebook", ".cbz") == ["PDF"]
    with pymupdf.open(run(cbz, "PDF")[0]) as d:
        assert d.page_count == 2
    svg = tmp_path / "Logo.svg"
    svg.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="200" height="100">'
                   '<rect width="200" height="100" fill="orange"/></svg>')
    from PIL import Image as I
    im = I.open(run(svg, "PNG")[0])
    assert im.size[0] >= 1000 and im.size[0] == 2 * im.size[1]
    assert I.open(run(svg, "JPG")[0]).mode == "RGB"
    assert run(svg, "PDF")[0].exists()
    # PDF extras
    assert "Chapter one" in run(pdf, "HTML")[0].read_text(encoding="utf-8")
    assert run(pdf, "SVG")[0].read_text(encoding="utf-8").startswith("<svg")


# ---------------------------------------------------------------- archives
def test_archive_round_trip(zipf):
    for target in ("TAR", "TAR.GZ", "TAR.XZ", "TAR.BZ2"):
        out = run(zipf, target)[0]
        assert out.name == ("Pack" + "." + target.lower()) and category(out) == "archive"
        with tarfile.open(out) as t:
            assert t.extractfile("dir/b.txt").read() == b"world" * 1000
    tgz = zipf.with_name("Pack.tar.gz")
    back = run(tgz, "ZIP")[0]
    with zipfile.ZipFile(back) as z:
        assert z.read("a.txt") == b"hello"
    assert "TAR.GZ" not in convert.targets_for("archive", ".tar.gz") and "ZIP" in convert.targets_for("archive", ".tar.gz")


def test_archive_rejects_unsafe_paths_and_garbage(tmp_path):
    evil = tmp_path / "evil.tar"
    with tarfile.open(evil, "w") as t:
        import io
        ti = tarfile.TarInfo("../../etc/passwd")
        ti.size = 1
        t.addfile(ti, io.BytesIO(b"x"))
    with pytest.raises(UserError, match="unsafe"):
        run(evil, "ZIP")
    assert not list(tmp_path.glob("evil*.zip")) and not list(tmp_path.glob(".~*"))
    bad = tmp_path / "bad.zip"
    bad.write_bytes(b"not a zip")
    with pytest.raises(UserError):
        run(bad, "TAR")


# ---------------------------------------------------------------- subtitles
@needs_ffmpeg
def test_subtitles(tmp_path):
    srt = tmp_path / "Subs.srt"
    srt.write_text("1\n00:00:01,000 --> 00:00:02,500\nHello\n\n2\n00:00:03,000 --> 00:00:04,000\nWorld\n")
    vtt = run(srt, "VTT")[0]
    assert vtt.read_text(encoding="utf-8").startswith("WEBVTT") and "Hello" in vtt.read_text(encoding="utf-8")
    back = run(vtt, "SRT")[0]
    assert "00:00:01,000 --> 00:00:02,500" in back.read_text(encoding="utf-8")
    assert "Dialogue" in run(srt, "ASS")[0].read_text(encoding="utf-8")


def test_existing_text_and_docx_conversions_unchanged(tmp_path):
    t = tmp_path / "n.txt"
    t.write_text("hello")
    assert convert.targets_for("text", ".txt") == ["PDF", "DOCX", "HTML"]
    docx = convert.convert(t, "text", "DOCX")[0]
    assert convert.targets_for("docx", ".docx")[:2] == ["PDF", "TXT"]
    assert convert.convert(docx, "docx", "TXT")[0].read_text(encoding="utf-8").strip() == "hello"
