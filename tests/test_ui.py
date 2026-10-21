"""The real tool windows and menu dispatch, driven headless (offscreen Qt) the way a user would:
pick options, press the main button, check the file that lands on disk."""
from __future__ import annotations

import io
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from PIL import Image
from PySide6.QtCore import QPointF, QRectF

from conftest import camera_exif, date_taken, exif_of, has_gps, icc_desc, photo, rgb
from dragshift import convert, tools

pytestmark = pytest.mark.usefixtures("qapp")


class Spy:
    """Stands in for the app's notify(): records what a window produced."""

    def __init__(self):
        self.calls = []

    def __call__(self, action, srcs, outs, extra=""):
        self.calls.append((action, outs, extra))

    @property
    def outs(self) -> list[Path]:
        return [Path(o) for _, outs, _ in self.calls for o in outs]


@pytest.fixture
def sideways_jpg(tmp_path, p3):
    """4000×3000 stored, orientation 6 → displays as 3000×4000 (a portrait phone photo)."""
    p = tmp_path / "IMG_2001.jpg"
    photo(4000, 3000).save(p, quality=93, exif=camera_exif(orientation=6), icc_profile=p3)
    return p


# ---------------------------------------------------------------- crop
def test_crop_full_frame_is_exact(sideways_jpg):
    spy = Spy()
    w = tools.CropWindow(str(sideways_jpg), spy)
    w.reset()
    w.apply()
    out = spy.outs[0]
    assert Image.open(out).size == (3000, 4000), "full-frame crop lost pixels in rounding"
    assert icc_desc(out) == "Display P3" and date_taken(out)
    assert exif_of(out).get(0x0112, 1) == 1


@pytest.mark.parametrize("ratio_name,ratio", [("1:1", 1.0), ("4:3", 4 / 3), ("16:9", 16 / 9)])
def test_crop_ratio_is_exact(sideways_jpg, ratio_name, ratio):
    spy = Spy()
    w = tools.CropWindow(str(sideways_jpg), spy)
    i = [n for n, _ in w.RATIOS].index(ratio_name)
    w.grp.button(i).click()
    w.apply()
    ow, oh = Image.open(spy.outs[0]).size
    assert abs(ow / oh - ratio) < 0.002 and ow <= 3000 and oh <= 4000


def test_crop_with_rotate(sideways_jpg):
    spy = Spy()
    w = tools.CropWindow(str(sideways_jpg), spy)
    w._rotate(90)
    w.apply()
    assert Image.open(spy.outs[0]).size == (4000, 3000)


# ---------------------------------------------------------------- redact
def test_redact_window_burns_box_at_full_resolution(sideways_jpg):
    spy = Spy()
    w = tools.RedactWindow(str(sideways_jpg), spy)
    k = w.k  # full px per preview px
    w.view.rects.append(QRectF(100, 100, 200, 50))  # preview coordinates
    w.save()
    out = spy.outs[0]
    a = rgb(out)
    y0, x0 = round(100 * k), round(100 * k)
    assert a[y0 + 10:y0 + round(50 * k) - 10, x0 + 10:x0 + round(200 * k) - 10].max() <= 10
    assert not has_gps(out) and not exif_of(out).get(0x010F)


def test_redact_without_boxes_writes_nothing(sideways_jpg):
    spy = Spy()
    tools.RedactWindow(str(sideways_jpg), spy).save()
    assert not spy.outs


# ---------------------------------------------------------------- resize / compress
def test_resize_window_batch(tmp_path):
    a, b = tmp_path / "a.jpg", tmp_path / "b.png"
    photo(2000, 1000).save(a, quality=95)
    photo(800, 1200).save(b)
    spy = Spy()
    w = tools.ResizeWindow([str(a), str(b)], spy)
    w.lock.setChecked(True)
    w.w.setValue(500)  # width 500, proportions kept, applied per file
    w.save()
    sizes = sorted(Image.open(o).size for o in spy.outs)
    assert sizes == [(500, 250), (500, 750)]


@pytest.mark.parametrize("target_idx", [2, 3, 4])  # 500 KB, 200 KB, 100 KB
def test_compress_window_target_size(sideways_jpg, target_idx):
    spy = Spy()
    w = tools.CompressWindow([str(sideways_jpg)], spy)
    w.mode.button(1).click()
    w.tgrp.button(target_idx).click()
    w.save()
    limit = w.TARGETS[target_idx][1]
    out = spy.outs[0]
    assert out.stat().st_size <= limit
    assert icc_desc(out) == "Display P3"
    assert Image.open(out).size[1] > Image.open(out).size[0], "portrait photo came out sideways"


def test_compress_window_quality_mode_webp(sideways_jpg):
    spy = Spy()
    w = tools.CompressWindow([str(sideways_jpg)], spy)
    w.webp.setChecked(True)
    w.save()
    out = spy.outs[0]
    assert out.suffix == ".webp" and out.stat().st_size < sideways_jpg.stat().st_size


# ---------------------------------------------------------------- edit / annotate / background
def test_edit_window_rotate_and_save(sideways_jpg):
    spy = Spy()
    w = tools.EditWindow(str(sideways_jpg), spy)
    w._rot(90)
    w.save()
    out = spy.outs[0]
    assert Image.open(out).size == (4000, 3000)
    assert icc_desc(out) == "Display P3" and date_taken(out)


def test_annotate_window_draws_at_full_resolution(tmp_path, p3):
    src = tmp_path / "shot.png"
    Image.new("RGB", (2400, 1600), (255, 255, 255)).save(src, icc_profile=p3)
    spy = Spy()
    w = tools.AnnotateWindow(str(src), spy)
    w.view.shapes.append(tools.Shape("Rect", "#ff0000", 6, [QPointF(50, 50), QPointF(300, 200)]))
    w.save()
    out = spy.outs[0]
    o = Image.open(out)
    assert o.size == (2400, 1600) and icc_desc(out) == "Display P3"
    assert (rgb(out)[:, :, 1] < 100).any(), "rectangle not drawn"


def test_add_background_window(tmp_path, p3):
    src = tmp_path / "shot.png"
    photo(800, 500).save(src, icc_profile=p3)
    spy = Spy()
    tools.AddBgWindow(str(src), spy).save()
    out = spy.outs[0]
    o = Image.open(out)
    assert out.name == "shot With Background.png" and o.width > 800 and icc_desc(out) == "Display P3"


# ---------------------------------------------------------------- metadata / protect
def test_metadata_window_buttons(sideways_jpg):
    spy = Spy()
    w = tools.MetadataWindow(str(sideways_jpg), spy)
    assert w.has_gps
    w.strip_location()
    w2 = tools.MetadataWindow(str(sideways_jpg), spy)
    w2.strip()
    loc, clean = spy.outs
    assert not has_gps(loc) and date_taken(loc)
    assert not has_gps(clean) and not date_taken(clean) and icc_desc(clean) == "Display P3"


def test_protect_window(tmp_path):
    d = pymupdf.open()
    d.new_page().insert_text((72, 72), "confidential")
    src = tmp_path / "c.pdf"
    d.save(src)
    spy = Spy()
    w = tools.ProtectWindow(str(src), spy)
    w.pw.setText("abc")
    w.pw2.setText("abc")
    assert not w.primary_btn.isEnabled(), "accepted a 3-character password"
    w.pw.setText("hunter22")
    w.pw2.setText("hunter2")
    assert not w.primary_btn.isEnabled(), "accepted mismatched passwords"
    w.pw2.setText("hunter22")
    assert w.primary_btn.isEnabled()
    w.save()
    out = pymupdf.open(spy.outs[0])
    assert out.needs_pass and out.authenticate("hunter22")


# ---------------------------------------------------------------- menu dispatch: every slice
class FakeApp:
    """DragShiftApp's real menu/dispatch code, with jobs run inline and windows captured."""
    from dragshift.app import DragShiftApp as _A
    build_rings = _A.build_rings
    run_item = _A.run_item

    def __init__(self):
        self.outs, self.windows, self.errors = [], [], []
        self.toasts = self

    def submit(self, action, srcs, fn, working="", extra=""):
        self.outs += [Path(o) for o in fn()]

    def open_window(self, w):
        self.windows.append(w)

    def notify(self, *a):
        pass

    def error(self, title, msg):  # toasts.error
        self.errors.append((title, msg))


def _sample_files(tmp_path, p3) -> list[Path]:
    files = []
    for ext, fmt in [("png", "PNG"), ("jpg", "JPEG"), ("heic", "HEIF"), ("webp", "WEBP"), ("gif", "GIF"),
                     ("tiff", "TIFF"), ("avif", "AVIF"), ("bmp", "BMP"), ("ico", "ICO")]:
        p = tmp_path / f"sample.{ext}"
        im = photo(256, 256) if ext == "ico" else photo(900, 600)
        im.save(p, fmt, **({"icc_profile": p3} if fmt not in ("GIF", "BMP", "ICO") else {}))
        files.append(p)
    d = pymupdf.open()
    for i in range(2):
        pg = d.new_page()
        pg.insert_text((72, 72), f"page {i}")
        buf = io.BytesIO()
        photo(300, 200).save(buf, "JPEG")
        pg.insert_image(pymupdf.Rect(72, 100, 372, 300), stream=buf.getvalue())
    d.save(tmp_path / "sample.pdf")
    files.append(tmp_path / "sample.pdf")
    doc = Document()
    doc.add_heading("Hello", 1)
    doc.add_paragraph("World")
    doc.save(tmp_path / "sample.docx")
    files.append(tmp_path / "sample.docx")
    for ext in ("txt", "md", "csv", "json"):
        (tmp_path / f"sample.{ext}").write_text("a,b\n1,2\n", "utf-8")
        files.append(tmp_path / f"sample.{ext}")
    return files


def test_every_menu_slice_for_every_file_type(tmp_path, p3, monkeypatch):
    for name in ("_word_to_pdf", "_pages_to_pdf", "_libreoffice_to_pdf"):
        monkeypatch.setattr(convert, name, lambda *a: False)  # don't drive Pages in this sweep
    tried = 0
    for f in _sample_files(tmp_path, p3):
        app = FakeApp()
        rings = app.build_rings([str(f)])
        assert rings.get("Convert"), f"no conversions offered for {f.name}"
        for ring in rings.values():
            for item in ring:
                before = len(app.outs) + len(app.windows)
                app.run_item(item, [str(f)])
                assert len(app.outs) + len(app.windows) > before, f"{f.name} → {item.key} did nothing"
                tried += 1
        assert not app.errors, (f.name, app.errors)
        for o in app.outs:
            assert o.exists() and o.stat().st_size > 0, o
        # every window opened from the menu can save with its defaults
        for w in app.windows:
            spy = Spy()
            w.notify = spy
            if isinstance(w, tools.ProtectWindow):
                w.pw.setText("password1")
                w.pw2.setText("password1")
            if isinstance(w, tools.RedactWindow):
                w.view.rects.append(QRectF(10, 10, 50, 50))
            if isinstance(w, tools.AnnotateWindow):  # saving with nothing drawn writes nothing
                w.view.shapes.append(tools.Shape("Arrow", "#ff0000", 4, [QPointF(5, 5), QPointF(60, 40)]))
            if isinstance(w, tools.CropWindow):
                w.apply()
            elif isinstance(w, tools.MetadataWindow):
                w.strip()
            else:
                w.save()
            assert spy.outs and all(o.exists() for o in spy.outs), (f.name, type(w).__name__)
    assert tried > 150


def test_multi_file_menu(tmp_path):
    imgs = []
    for i in range(3):
        p = tmp_path / f"p{i}.jpg"
        photo(400, 300, seed=i).save(p)
        imgs.append(str(p))
    app = FakeApp()
    rings = app.build_rings(imgs)
    pdf = [it for it in rings["Convert"] if it.key == "convert:PDF"][0]
    app.run_item(pdf, imgs)
    assert len(app.outs) == 1 and pymupdf.open(app.outs[0]).page_count == 3
    app = FakeApp()
    app.run_item([it for it in rings["Convert"] if it.key == "convert:WEBP"][0], imgs)
    assert len(app.outs) == 3


# ---------------------------------------------------------------- background job + errors
def _half_jpeg() -> bytes:
    """A real photo cut off halfway, like an interrupted download."""
    buf = io.BytesIO()
    photo(800, 600).save(buf, "JPEG", quality=90)
    return buf.getvalue()[: len(buf.getvalue()) // 2]


@pytest.mark.parametrize("make,expect", [
    (lambda p: p.write_bytes(b""), "doesn't look like an image"),
    (lambda p: p.write_bytes(_half_jpeg()), "damaged"),
])
def test_job_reports_friendly_errors(tmp_path, make, expect):
    from dragshift.app import Job
    src = tmp_path / "bad.jpg"
    make(src)
    job = Job("convert:PNG", [str(src)], lambda: convert.convert(src, "image", "PNG"))
    got = {}
    job.sig.failed.connect(lambda m: got.setdefault("msg", m))
    job.sig.done.connect(lambda r: got.setdefault("done", r))
    job.run()
    assert expect in got.get("msg", ""), got
    assert not list(tmp_path.glob(".~*")), "temp file left behind"


def test_job_permission_error(tmp_path):
    from dragshift.app import Job
    src = tmp_path / "a.png"
    photo(64, 64).save(src)

    def locked():  # a read-only folder; raised directly because Windows and root ignore chmod
        raise PermissionError(13, "Permission denied", str(src.with_suffix(".jpg")))
    job = Job("convert:JPG", [str(src)], locked)
    got = {}
    job.sig.failed.connect(lambda m: got.setdefault("msg", m))
    job.run()
    assert "No permission" in got["msg"]


def test_job_success_report(tmp_path):
    from dragshift.app import Job
    src = tmp_path / "a.png"
    photo(1200, 800).save(src)
    job = Job("convert:JPG", [str(src)], lambda: convert.convert(src, "image", "JPG"))
    got = {}
    job.sig.done.connect(lambda r: got.setdefault("r", r))
    job.run()
    assert got["r"].outs and got["r"].outs[0].exists()
